from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import hashlib, json, re, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('D:/Codex/GRL研究')
SECTIONS=ROOT/'outputs/frame_paper_sections_20261017'
OUT=ROOT/'outputs/frame_cn_merged_20261017_v1'
OUT.mkdir(parents=True,exist_ok=True)
NAME='FRAME_ICAPS2027_中文合稿_V1'
assert not (OUT/'交付完成.json').exists(),'Preserve delivered outputs; create a new version.'
sources={
 'abstract':'01_摘要/FRAME_中文摘要_用户确认稿_V3.md',
 'introduction':'02_引言/FRAME_中文引言_表达修订_V3.md',
 'methods':'04_方法/FRAME_中文方法_文字修订_V2.md',
 'experiments':'05_实验/FRAME_中文实验_文字初稿_V1.md',
 'related_work':'03_相关工作/FRAME_中文相关工作_引文修订_V2.md',
 'references':'07_参考文献/FRAME_参考文献_核定稿_V1.md'}
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
raw={k:(SECTIONS/v).read_text(encoding='utf-8') for k,v in sources.items()}
abstract=raw['abstract'].split('\n\n',1)[1].split('\n---',1)[0].strip()

def section(k,n,title):
    lines=raw[k].strip().splitlines()[1:]
    out=[f'# {n} {title}','']
    h2=0
    for line in lines:
        if line.startswith('## '):
            h2+=1;label=re.sub(r'^\d+(?:\.\d+)*\s+','',line[3:])
            line=f'## {n}.{h2} {label}'
        elif line.startswith('### '):
            label=line[4:];m=re.match(r'\d+\.(\d+)\s+(.*)',label)
            assert m,label
            line=f'### {n}.{h2}.{m.group(1)} {m.group(2)}'
        out.append(line)
    return '\n'.join(out).strip()

conclusion='''# 4 结论

本文提出反馈驱动的在线电网拓扑搜索框架FRAME，通过前缀筛查、分批核查和同站补查分配候选评估计算。框架根据已经获得的仿真结果调整搜索深度，并利用变电站信息选择补查对象，可直接接入现有控制器，无需额外训练搜索控制策略。在6个生成时序场景上，相对完整候选搜索，FRAME的查询量和完整控制时间宏平均分别减少65.10%和14.98%，全部场景达到预设运行质量标准。在13个既有场景中，FRAME以更少查询取得了比固定512策略更高的质量达标率。与学习式元推理适配对照相比，FRAME也表现出小幅平均计算优势。实验表明，利用搜索反馈调整评估数量和范围，可以减少拓扑控制中的仿真开销，并提高完整控制器的运行效率。'''
limitations='''# 6 局限

本文评价基于同一公开电网的两组时序场景，采用既定候选库和候选顺序。现有结果反映了这一实验设置下的表现，其他电网规模、运行条件和候选排序仍需进一步检验。6个生成场景在首次评价后继续用于对照与机制分析，后续比较属于同场景分析。

FRAME通过有限查询选择动作，其运行效果取决于候选序列与反馈质量。同站补查在部分状态中能够找到更好的动作，仍可能遗漏其他变电站的高收益候选。本文的质量标准用于比较已测场景中的闭环运行效果，长期安全性还需要结合电网运行约束和专门的安全分析。会议方法对照采用面向电网任务的适配实现，所报告差距对应本次适配配置。'''

parts=['# FRAME反馈驱动的在线电网拓扑搜索','## 摘要',abstract,
       section('introduction',1,'引言'),section('methods',2,'方法'),
       section('experiments',3,'实验'),conclusion,section('related_work',5,'相关工作'),
       limitations,raw['references'].strip()]
merged='\n\n'.join(parts)+'\n'
# Unify citation presentation without changing the approved abstract or scientific results.
merged=re.sub(r'\[([^\]]+)\]\(https?://[^\s)]+\)',lambda m:'（'+m.group(1)+'）',merged)
merged=merged.replace('实验采用Grid2Op公开电网环境l2rpn_idf_2023，','实验采用Grid2Op公开电网环境l2rpn_idf_2023（Donnot，2020），')
merged=merged.replace('基于（Bhatia等提出的ICAPS 2022元推理方法）构建','基于Bhatia等（2022）提出的元推理方法构建')
assert abstract in merged
assert '撤稿' not in merged and '周毅' not in merged
assert '![' not in merged and not re.search(r'^\s*\|',merged,re.M)
(OUT/(NAME+'.md')).write_text(merged,encoding='utf-8')
(OUT/'04_结论_合稿新增.md').write_text(conclusion+'\n',encoding='utf-8')
(OUT/'06_局限_合稿新增.md').write_text(limitations+'\n',encoding='utf-8')
shutil.copy2(SECTIONS/'07_参考文献/references.bib',OUT/'references.bib')

def mn(tag):return OxmlElement('m:'+tag)
def mr(text,script=False):
    e=mn('r');pr=mn('rPr')
    if script:
        a=mn('scr');a.set(qn('m:val'),'script');pr.append(a)
    else:
        a=mn('sty');a.set(qn('m:val'),'p');pr.append(a)
    e.append(pr);t=mn('t');t.text=text;e.append(t);return e
def seq(nodes):
    e=mn('e')
    for x in nodes:e.append(mr(x) if isinstance(x,str) else x)
    return e
def sub(base,index):
    el=mn('sSub');el.append(seq([base]));idx=mn('sub')
    for x in index if isinstance(index,list) else [index]:idx.append(mr(x) if isinstance(x,str) else x)
    el.append(idx);return el
def sup(base,index):
    el=mn('sSup');el.append(seq([base]));idx=mn('sup')
    for x in index if isinstance(index,list) else [index]:idx.append(mr(x) if isinstance(x,str) else x)
    el.append(idx);return el
def bar(base):
    e=mn('bar');pr=mn('barPr');pos=mn('pos');pos.set(qn('m:val'),'top');pr.append(pos)
    e.extend([pr,seq([base])]);return e

class MathParser:
    symbols={'rho':'ρ','pi':'π','tau':'τ','oplus':'⊕','rightarrow':'→','ldots':'…',
             'varnothing':'∅','ne':'≠','neq':'≠','in':'∈','setminus':'∖','min':'min','max':'max'}
    def __init__(self,s):self.s=s;self.i=0
    def atom(self):
        s=self.s
        while self.i<len(s) and s[self.i].isspace():self.i+=1
        assert self.i<len(s),s
        c=s[self.i];self.i+=1
        if c=='{':
            ns=self.parse('}');group=mn('box');group.append(seq(ns));return group
        if c=='\\':
            m=re.match(r'[A-Za-z]+',s[self.i:])
            if not m:
                c=s[self.i];self.i+=1;return mr(c)
            cmd=m.group();self.i+=len(cmd)
            if cmd in ('mathcal','mathrm','operatorname'):
                word=re.match(r'\{([A-Za-z]+)\}',s[self.i:])
                if word:
                    self.i+=len(word.group());return mr(word.group(1),script=cmd=='mathcal')
            if cmd in ('mathcal','mathrm','operatorname'):
                node=self.atom()
                if cmd=='mathcal':
                    for run in node.iter(qn('m:r')):
                        pr=run.find(qn('m:rPr'));scr=mn('scr');scr.set(qn('m:val'),'script');pr.append(scr)
                return node
            if cmd=='bar':return bar(self.atom())
            assert cmd in self.symbols,cmd
            return mr(self.symbols[cmd])
        return mr(c)
    def parse(self,stop=None):
        out=[]
        while self.i<len(self.s):
            if stop and self.s[self.i]==stop:self.i+=1;break
            if self.s[self.i].isspace():self.i+=1;continue
            node=self.atom()
            while self.i<len(self.s) and self.s[self.i] in '_^':
                op=self.s[self.i];self.i+=1;arg=self.atom()
                node=sub(node,[arg]) if op=='_' else sup(node,[arg])
            out.append(node)
        return out

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(.85)
sec.footer_distance=Inches(.35)
def stylefont(name,sz,ea='宋体',latin='Times New Roman'):
    st=doc.styles[name];st.font.name=latin;st.font.size=Pt(sz);st.font.color.rgb=RGBColor(0,0,0)
    rp=st.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for key,val in [('ascii',latin),('hAnsi',latin),('eastAsia',ea)]:rf.set(qn('w:'+key),val)
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    pf=st.paragraph_format;pf.widow_control=True
    return st
normal=stylefont('Normal',11.5);normal.paragraph_format.line_spacing=1.25
normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',19),('Heading 1',15),('Heading 2',12.5),('Heading 3',11.5)]:
    st=stylefont(name,size,'黑体');st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(5)
doc.styles['Title'].paragraph_format.space_before=Pt(0)
stylefont('No Spacing',10,'宋体','Consolas')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='FRAME反馈驱动的在线电网拓扑搜索'
doc.core_properties.author='';doc.core_properties.subject='在线电网拓扑搜索'

def inline(p,text):
    for tok in re.split(r'(\$[^$]+\$|(?<!\w)\*[^*\n]+\*(?!\w)|`[^`]+`)',text):
        if not tok:continue
        if tok.startswith('$'):
            om=mn('oMath')
            for node in MathParser(tok[1:-1]).parse():om.append(node)
            p._p.append(om)
        else:
            italic=bool(re.fullmatch(r'\*[^*\n]+\*',tok));mono=tok.startswith('`')
            run=p.add_run(tok[1:-1] if italic or mono else tok);run.italic=italic
            if mono:run.font.name='Consolas';run.font.size=Pt(10)

def display(number):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(8)
    om=mn('oMath')
    if number==1:
        ns=['0 < ',sub('ρ','s'),'(c) < τ(s),   τ(s) = min{',sub('ρ','cur'),'(s), ',bar('ρ'),'}']
    elif number==2:
        lim=mn('limLow');lim.append(seq(['arg max']));idx=mn('lim')
        idx.extend([mr('c ∈ '),sub('E','valid')]);lim.append(idx)
        ns=[sup('c','*'),' = ',lim,' ',sub('R','s'),'(c)']
    else:ns=[sub(mr('L',True),'v'),' = {π(a) : a ∈ ',mr('A',True),', g(a) = v} ∖ E']
    for node in ns:om.append(mr(node) if isinstance(node,str) else node)
    p._p.append(om);p.add_run('    （'+str(number)+'）')

lines=merged.splitlines();in_code=False;in_math=False;mathnum=0;ref=False;codeparas=[]
for line in lines:
    if line.startswith('```'):
        if in_code and codeparas:
            # Keep each complete algorithm together on a page where possible.
            for p in codeparas[:-1]:p.paragraph_format.keep_with_next=True
        in_code=not in_code;continue
    if in_code:
        p=doc.add_paragraph(style='No Spacing');p.paragraph_format.line_spacing=1.03
        p.paragraph_format.space_after=Pt(1);p.paragraph_format.left_indent=Inches(.08)
        p.add_run(line if line else ' ');codeparas.append(p);continue
    if line.strip()=='\\[':in_math=True;mathnum+=1;display(mathnum);continue
    if line.strip()=='\\]':in_math=False;continue
    if in_math or not line.strip():continue
    if line.startswith('# FRAME'):
        p=doc.add_paragraph(line[2:],style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;continue
    if line=='## 摘要':doc.add_heading('摘要',level=1);continue
    if line.startswith('# '):
        label=line[2:];ref=label=='参考文献';doc.add_heading(label,level=1);continue
    if line.startswith('## '):doc.add_heading(line[3:],level=2);continue
    if line.startswith('### '):doc.add_heading(line[4:],level=3);continue
    p=doc.add_paragraph()
    if ref:
        p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.line_spacing=1.1
        p.paragraph_format.keep_together=True
    else:p.paragraph_format.first_line_indent=Pt(23)
    inline(p,line)
    if ref:
        for run in p.runs:run.font.size=Pt(11)

doc.save(OUT/(NAME+'.docx'))
receipt={'utc':datetime.now(timezone.utc).isoformat(),
 'source_sections':{k:{'path':v,'sha256':sha(SECTIONS/v)} for k,v in sources.items()},
 'approved_abstract_verbatim':abstract in merged,
 'newly_composed_sections':['4 结论','6 局限'],
 'integration_changes':['统一章节与子节编号','统一正文引文外观','为Grid2Op补入已核定Donnot2020引用'],
 'scientific_data_changed':False,'figures':0,'tables':0,
 'editable_equations':'native OMML; three display equations and inline symbols',
 'files':{n:sha(OUT/n) for n in [NAME+'.md',NAME+'.docx','references.bib']}}
(OUT/'合稿来源与修改说明.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx':str(OUT/(NAME+'.docx')),'paragraphs':len(doc.paragraphs),
                  'inline_and_display_math':len(doc._element.findall('.//'+qn('m:oMath')))},ensure_ascii=False))
