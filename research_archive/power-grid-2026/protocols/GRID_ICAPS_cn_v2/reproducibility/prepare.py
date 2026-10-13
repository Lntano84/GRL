from pathlib import Path
import re,json,hashlib,difflib,shutil
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
HERE=Path(__file__).resolve().parent
OUT=ROOT/'outputs/grid_icaps_cn_v2'
OUT.mkdir(exist_ok=True)
user=ROOT/'outputs/grid_icaps_cn_v1/ICAPS2027_中文完整初稿.docx'
userhash=hashlib.sha256(user.read_bytes()).hexdigest()
receipt=OUT/'source_receipt.json'
if receipt.exists():
    assert json.loads(receipt.read_text(encoding='utf-8'))['user_docx_sha256']==userhash
else:
    receipt.write_text(json.dumps({'user_docx':str(user),'user_docx_sha256':userhash,'policy':'Preserve user document unchanged; write separate full rewrite.'},ensure_ascii=False,indent=2),encoding='utf-8')
    (HERE/'user_source_snapshot').mkdir(exist_ok=True)
    shutil.copy2(user,HERE/'user_source_snapshot'/user.name)
old=(ROOT/'work/grid_cn_draft/manuscript_zh.md').read_text(encoding='utf-8')
src=(HERE/'manuscript_zh.md').read_text(encoding='utf-8')
tables=re.findall(r'(?m)^\|.*(?:\n\|.*)*',old)
for i,t in enumerate(tables):src=src.replace(f'@@TABLE{i}@@',t)
refs=old.split('## 参考文献\n',1)[1].strip()
refs=refs.replace(' 本文方法讨论对应这一具体版本；后续改题版本不计为另一篇工作。','')
src=src.replace('@@REFERENCES@@',refs)
assert '@@' not in src
assert not any(x in src for x in ['不是','而不是','不意味着','不能声称','尚不能','不代表'])
(HERE/'manuscript_ready.md').write_text(src,encoding='utf-8')
old_doc=ROOT/'work/GRL-results-integration/research_archive/power-grid-2026/protocols/GRID_ICAPS_cn_v1/ICAPS2027_中文完整初稿.docx'
if old_doc.exists():
    a=[p.text for p in Document(old_doc).paragraphs if p.text]
    b=[p.text for p in Document(user).paragraphs if p.text]
    (HERE/'user_edits.diff').write_text('\n'.join(difflib.unified_diff(a,b,fromfile='generated_v1',tofile='user_edited_v1',lineterm='')),encoding='utf-8')
b=(ROOT/'work/grid_cn_draft/build.py').read_text(encoding='utf-8')
b=b.replace("outputs/grid_icaps_cn_v1","outputs/grid_icaps_cn_v2").replace("SRC=HERE/'manuscript_zh.md'","SRC=HERE/'manuscript_ready.md'")
b=b.replace('ICAPS2027_中文完整初稿','ICAPS2027_中文重写稿_v2')
b=b.replace('面向在线电网控制的质量与计算权衡','面向在线电网控制的高效候选评估')
b=b.replace("doc.styles['Normal'].paragraph_format.line_spacing=1.2","doc.styles['Normal'].paragraph_format.line_spacing=1.15")
b=b.replace("        p=doc.add_paragraph('面向在线电网控制的高效候选评估',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER\n",'')
b=b.replace('只使用公开状态与实际查询反馈 排序器可替换 各方法共有基础控制与缓存','输入当前电网状态 输出经过预测验证的拓扑提案')
b=b.replace('补全只在扩展停止分支执行 不递归补新候选的站 没有可接纳项则继续直到池耗尽','全局渐进评估定位有效变电站 局部补查完成同站配置比较')
b=b.replace('单状态选择键不降低 不代表全局最优 长期安全或硬实时保证','复用已查询反馈 保留最佳候选 每个代表最多评估一次')
b=b.replace('图 1 共同控制器中的渐进搜索与条件补全','图 1 渐进搜索与变电站补全框架')
b=b.replace("len(english_heads)==len(chinese_heads)==11","len(chinese_heads)==8")
b=b.replace("assert len(re.findall(r'^### ',source,re.M))==len(re.findall(r'^### ',en.read_text(encoding='utf-8'),re.M))","assert re.findall(r'^## (.*)',source,re.M)==['摘要','1 引言','2 方法','3 实验','4 结论','5 相关工作','6 局限','参考文献']")
b=b.replace("in_refs=(s[3:]=='参考文献')","in_refs=(s[3:]=='参考文献')")
b=b.replace('full_sections=11','full_sections=8')
b=b.replace('Full translation and document construction. Table/equation structure verified; semantic translation reviewed by the authoring agent, not independent peer review. Render QA required before delivery.','Full Chinese rewrite using user edits as style guidance. Same three numerical tables and five equations; source user document preserved. Render QA required.')
b=b.replace("translation_check.json","rewrite_check.json")
(HERE/'build.py').write_text(b,encoding='utf-8')
r=(ROOT/'work/grid_cn_draft/render.py').read_text(encoding='utf-8')
r=r.replace('outputs/grid_icaps_cn_v1','outputs/grid_icaps_cn_v2').replace('ICAPS2027_中文完整初稿','ICAPS2027_中文重写稿_v2')
r=r.replace("Path(__file__).resolve().parent/'render_dependency/receipt.json'","ROOT/'work/grid_cn_draft/render_dependency/receipt.json'")
r=r.replace("is_relative_to(Path(__file__).resolve().parent)","is_relative_to(ROOT/'work/grid_cn_draft')")
(HERE/'render.py').write_text(r,encoding='utf-8')
print(json.dumps({'user_preserved':True,'sections':re.findall(r'^## (.*)',src,re.M),'han_chars':len(re.findall('[\u4e00-\u9fff]',src))},ensure_ascii=False))
