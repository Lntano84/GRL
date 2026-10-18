"""Reorganize audited evidence into a contribution-focused Chinese manuscript."""
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import csv
import hashlib
import json
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('D:/Codex/GRL研究')
OUT = ROOT / 'outputs/frame_cn_v3_20261018'
NAME = 'FRAME_ICAPS2027_中文合稿_V3'
OLD = ROOT / 'outputs/grid89_review/manuscript_v2'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
read = lambda p: json.loads(p.read_text(encoding='utf-8'))


def main():
    assert not (OUT / '交付完成.json').exists(), 'Preserve delivered V3'
    prior = OLD / 'FRAME_ICAPS2027_中文合稿_V2.md'
    old_text = prior.read_text(encoding='utf-8')
    refs = old_text.split('# 参考文献', 1)[1].strip()
    path = OUT / (NAME + '.md')
    text = path.read_text(encoding='utf-8')
    if not text.split('# 参考文献', 1)[1].strip():
        text = text.rstrip() + '\n\n' + refs + '\n'
        path.write_text(text, encoding='utf-8')
    shutil.copy2(OLD / 'references.bib', OUT / 'references.bib')
    # Keep exactly the previously verified Word styles, editable equations and algorithm layout.
    source = (ROOT / 'work/grid89_review/source_v1/build_reference.py').read_text(encoding='utf-8')
    renderer = source[source.index('def mn(tag)'):source.index("receipt={'utc'")]
    namespace = dict(globals(), merged=text)
    exec(compile(renderer, str(Path(__file__).resolve()), 'exec'), namespace)
    doc = namespace['doc']
    assert len(doc.tables) == len(doc.inline_shapes) == 0
    review = read(ROOT / 'outputs/grid89_review/review.json')
    decision = read(ROOT / 'outputs/grid89_review/decision.json')
    primary = next(c for c in decision['comparisons'] if c['phase'] == 'sealed' and c['method'] == 'FRAME_GUARD' and c['baseline'] == 'FULL')
    assert decision['sealed_frame_quality_dates'] == 6
    assert f"{primary['macro_controller_reduction'] * 100:.2f}" == '14.30'
    assert f"{primary['macro_query_reduction'] * 100:.2f}" == '60.73'
    assert primary['same_horizon_all_dates'] and primary['both_quality_all_dates']
    supplementary = '''# FRAME中文稿V3 补充材料

## A 实验配置与结果来源

本文统一使用FRAME表示最终规则，代码名称FRAME_GUARD。最终规则与work/grid70/policies.py一致，GRID89默认版本在137个既有完整状态上核对逐查询与最终选择，并在保存数据审计中复现既有轨迹。初始前缀已经达到阈值时，后续核查未改变前缀动作就不触发同站补查，这一条件已包含在算法1中。

正文主评价为G2：主种子20261017，六个时序ID为2036-01-07_10、2036-03-03_10、2036-05-05_10、2036-07-07_10、2036-09-01_10、2036-11-03_10。FULL、FRAME_GUARD、GLOBAL_MATCH与FIXED_512各运行三次，合计72条，来自outputs/grid89_review/sealed。主结果使用同批控制时间，逐日期先平均重复，再计算相对比例的日期宏平均。

D为13个2035年既有场景。FRAME与固定512的同批两次反向计时来自GRID88，共52条；质量达标13/13和11/13，查询减少33.57%，控制时间减少3.88%。固定256的13条质量回归来自GRID87，质量9/13；固定256未与旧控制时间混算提速。逐场景清单由outputs/grid88/design.json与outputs/grid87/design.json保存。

G1为主种子20261016生成的六个时序，日期与G2相同、时序ID后缀为_9。原FULL同批比较来自GRID81：FRAME与FULL均为6/6质量达标，均完成4/6整周，另外两例同一步终止；查询减少65.10%，控制时间减少14.98%。这组结果移出摘要和主结果，作为原有证据保留。

正文META比较来自GRID84：G1六个场景，成本0，种子0、1、2、3、4，末端权重，各两次反向计时。五个种子均为6/6质量达标，FRAME控制时间平均减少2.23%。逐种子控制时间降幅为7.55%、−0.30%、0.21%、1.32%、2.36%（最终数值以GRID84 decision.json逐种子字段为准）。META包含初始委托FRAME的动作，按五种子分别有12.5%、90.0%、76.7%、43.3%、56.7%的搜索使用该动作。

GRID89的META复查另使用单次计时，保留零成本五种子，并比较成本0.0025的五种子。只对非零成本种子2、3、4各新增一次训练，0、1复用既有权重。该批零成本平均时间差1.15%，非零成本平均时间差1.51%；非零成本五种子的质量达标数为5、6、6、6、5。两批的计时次数与结果分别保留，正文主要META比较采用已有两重复批次。

## B 同状态机制证据

同状态分析使用FULL实际访问状态中已揭示的候选反馈，比较NO_LOCAL、FRAME_GUARD、GLOBAL_MATCH及固定预算规则，不新增仿真。追加预算由同站剩余代表数确定，局部补查与全局扩展在同一状态的总查询量相同。缺少可用候选或无法计算奖励差的状态保留为缺失，不记为零。

D与G1来自GRID85；G2来自GRID89。D、G1、G2中FRAME的日期平均奖励遗憾分别为0.000726、0.000871、0.001153；等预算全局扩展分别为0.002156、0.000089、0.001823。正文保留G1遗憾反转及其关键案例，动作一致率仅作补充指标。

G1反向案例为2036-05-05_9的call000、第406步。全局阶段256次查询后，最佳候选62归属48号变电站；198次追加同站查询选候选748，同属48号站；全局扩展选候选285，归属64号站。候选的全局顺序位置分别为63、749、286，奖励分别约1.000857、1.001464、1.013823。分析使用outputs/grid81/runs/r0__2036-05-05_9__FULL/call000.json及GRID85等预算记录。

## C 候选顺序与参数检查

以下逐项结果来自冻结GRID89方案；未按结果替换默认顺序或参数。顺序比较为G1六场景、五种顺序各FULL与FRAME一次回放。默认教师频次、随机11、随机23、随机37与线路负载顺序的质量达标数依次为6、6、5、5、6；匹配活跃时域查询降幅依次为65.10%、58.58%、53.45%、58.78%、64.88%。仅一次计时，质量与查询为主要观察。

参数检查固定D中的2035-03-05_0、2035-06-25_8、2035-09-03_4、2035-10-01_5，共四场景，各配置一次回放。默认p128、h64、q_min256、L2、分支阈值0.9，质量4/4，查询9108次。

前缀p=64与256：质量4/4与3/4，查询9151与12119次。

批次h=32与128：质量4/4与4/4，查询8772与13904次。

最低核查量q_min=128与512：质量4/4与3/4，查询9108与14620次。

停滞批数L=1与3：质量4/4与4/4，查询8834与10972次。

分支阈值0.85与0.95：质量3/4与3/4，查询9491与9555次。

## D META训练与接口

一条训练记录是一个已保存的拓扑搜索状态，包含该状态的候选顺序、公开特征及已有候选反馈，不是一周轨迹。按日期家庭划分得到129条训练记录、22条校准记录；训练回合从训练记录中抽取一个搜索状态，以候选前缀推进元决策。网络为25—64—32—3的ReLU网络，训练30000回合，采用预设末端检查点。

动作包括返回当前候选、续查下一批64个候选、在初始阶段委托完整FRAME。25维输入涵盖查询比例、候选数量、可用候选标记、最佳奖励与负载率、最新改善、停滞程度、可用比例、质量均值与标准差、前两候选差、负载率限值及电网公开状态统计，精确字段列于GRID89 evidence/meta_interface_and_split.json。

质量q为当前最佳可用奖励R除以2后截断到[0,1]，无可用候选时为零；查询比例z=Q/M，beta=log(1+lambda)，效用U=q−[exp(beta·z)−1]，反馈为效用增量。正常候选奖励采用2减去最大负载率，严格可用性要求负载率大于零。成本0.0025对应查询完整候选库时显式惩罚上限0.0025。原预设三成本为0.25、0.0025、0，原校准未满足质量匹配条件，按已披露的备选规则选成本0；该事实在完整记录中保留。

学习率0.0001，折扣0.99，批量256，回放容量1000000，目标保留率0.995，探索最终值0.1。其余训练循环、状态变换、掩码与推理规则见work/grid77及GRID84、GRID89封存文件。全部学习曲线、初始与末端权重随记录保存，未在本次文稿重组中重新训练。

## E 复现资料与测量边界

LJN上游固定提交ca0637eab9f098be7f206ed0e46a3900cd4deec0用于追溯原始资源；本研究控制器与搜索实现按各轮code_freeze.json及data_freeze.json精确封存。候选以原始编号保留动作归属，组合动作按基础控制器的等价规则合并，候选比较按float32奖励及较小编号破并列。教师库原始采集日期覆盖尚未独立核实。

G2生成使用与G1相同官方Chronix2Grid wrapper及静态电网配置。逐日期种子、文件收据、完整生成参数和输入哈希见GRID89 generation.json、design.json、data_freeze.json及六份generation_receipt。数值环境、缓存与控制时间边界均沿用冻结实现；首次构建与当前字段验证按原规则计入控制时间，结果保存和物理推进另计。

G2各日期三次控制时间与全部/搜索步骤P50、P95、P99见GRID89 per_run.csv和per_date.csv；质量及查询也逐日期保留。审计核对246条回放、455427个控制步、1539次搜索和793809次候选查询，包括保存向量、合法反馈、费用、逻辑重放、48条重复轨迹及46条既有轨迹对应，不执行独立AC重求解。

本次只重组论文，所有实验、冻结判据和历史较差结果保留。GRID89完整结果已归档到research_archive/power-grid-2026/protocols/GRID89_review/complete；本文V3与V2及作者修改稿分别保存。
'''
    # Populate paired META numbers from the audited record rather than handwritten rounding.
    old_decision = read(ROOT / 'outputs/grid84/decision.json')
    reductions = '、'.join(f"{r['macro_controller_reduction'] * 100:.2f}%" for r in old_decision['decisions'])
    supplementary = re.sub(r'逐种子控制时间降幅为.*?。META包含', '逐种子控制时间降幅为' + reductions + '。META包含', supplementary)
    (OUT / '补充材料_V3.md').write_text(supplementary, encoding='utf-8')
    editorial = '''# V3贡献主线与材料取舍

核心主张：根据已有仿真反馈分配全局查询深度，再由当前最佳动作引导同站补查，减少在线拓扑搜索计算并保持已测闭环运行效果。

主评价采用固定方法后生成的G2：60.73%查询减少、14.30%完整控制时间减少；六日期质量全部达标，五个整周完成，一个与FULL同一步提前终止。摘要、引言贡献和结论统一使用该组。

机制证据按问题组织。查询深度：固定256/512回归、前缀已达阈值后仍改变动作的记录、相同状态的晚发现动作案例。查询位置：等预算补查替换、奖励遗憾、同站与远站反向案例、G2接近查询量下的闭环质量差异。计算来源：互斥耗时分解和搜索步骤延迟。

META保留GRID84五种子两重复作为次要竞争力比较，明确成本0、使用G1、包含FRAME委托。单次计时的1.15%/1.51%与非零成本逐项结果移入补充材料，不再与2.23%并列形成主结果。

从正文移出的内容：G1对FULL的旧主结果单独段落、换序与参数逐项罗列、META效用与训练超参数细节、重复的算法不变量与在线流程说明。完整结果继续保留，正文仍交代顺序/参数质量失败、G1同站收益反转和全步骤P95上升，避免主张失去适用条件。

数学和命名修订：统一FRAME与FRAME_GUARD；代表属于原始动作集合并保留原始编号，从而明确定义g(c)；原始别名经pi映射参与同站补查。说明float32奖励和编号破并列；明确两个停止分支及“初始达标但候选未变则不补查”。

本轮未新增实验、训练、模型或统计选型。V1、V2、作者修改稿与全部原始记录不覆盖。正文保持中文阅读版且不新增图表，后续配图与英文投稿排版分别处理。
'''
    (OUT / '贡献主线与修改说明.md').write_text(editorial, encoding='utf-8')
    source_paths = [prior, OLD / 'FRAME_ICAPS2027_中文合稿_V2.docx', ROOT / 'outputs/grid89_review/review.json', ROOT / 'outputs/grid89_review/decision.json', ROOT / 'outputs/grid84/decision.json', ROOT / 'outputs/grid84/behavior.json', ROOT / 'outputs/grid85_mechanism/equal_budget_pairs.csv', ROOT / 'outputs/grid88/GRID88_report.md']
    receipt = dict(utc=datetime.now(timezone.utc).isoformat(), version='V3 contribution-focused reorganization',
                   input_sha256={p.relative_to(ROOT).as_posix(): sha(p) for p in source_paths},
                   source_v2_preserved=True, new_experiments=0, new_fits=0, scientific_results_changed=False,
                   abstract_changed='Use audited G2 as the single primary result cohort; META remains secondary in text.',
                   figures=0, tables=0, editable_math_objects=len(doc._element.findall('.//' + qn('m:oMath'))),
                   paragraphs=len(doc.paragraphs), visual_review_pending=True,
                   files={p.name: sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in ['.md', '.docx', '.bib']})
    (OUT / 'source_receipt.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(dict(document=str(OUT / (NAME + '.docx')), paragraphs=len(doc.paragraphs), math=receipt['editable_math_objects']), ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
